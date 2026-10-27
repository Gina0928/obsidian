#!/usr/bin/env python3
"""
Java bridge for the excel-validate-stats skill.

This module lets the skill PREFER the original Spring Boot JAR
(`demo-0.0.1-SNAPSHOT.jar`) when it is available, and lets the Python
fallback (validate_merge.py / compare_and_validate.py) take over only when
the JAR cannot be found. Discovery is JAR-first and safe: explicit --jar >
cached path > source-project target/ walk > known locations (home
subdirectories are never scanned, avoiding sandbox blocks). The JAR exposes
two modes:

    validate <basePath> <sourceFile>          -> <source>_重复行合并.xlsx
    compare  <basePath> <sourceFile> <cmpFile> -> <source>_差异.xlsx

Both modes write their output into <basePath>. To keep the user's source
directories clean, we copy the inputs into a temporary work dir, run the
JAR there, then move the resulting file next to the source and delete the
temp dir.

Java itself does not understand Git-Bash style paths (/d/...), so every path
handed to `java` is normalized to a native Windows absolute path (backslashes)
via inv_io.normalize_path + str(Path).

Requires `java` on PATH. If Java is missing or the JAR fails, callers decide
whether to fail loud (explicit --jar) or degrade to Python (auto-detect).
"""

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from inv_io import normalize_path, insert_suffix, file_fingerprint, verify_inputs_untouched  # noqa: E402


JAR_NAME = "demo-0.0.1-SNAPSHOT.jar"

# User-level cache of the last-known-good JAR path. Lives OUTSIDE the skill
# directory (~/.workbuddy/...) so it survives skill updates, and makes every
# run after the first one deterministic and instant.
_CACHE_FILE = Path.home() / ".workbuddy" / "excel-validate-stats.jar.txt"


def _read_cache():
    """Return the cached JAR path if it still exists, else None."""
    try:
        p = Path(_CACHE_FILE.read_text(encoding="utf-8").strip())
        return p.resolve() if p.exists() else None
    except Exception:
        return None


def _write_cache(jar_path):
    try:
        _CACHE_FILE.parent.mkdir(parents=True, exist_ok=True)
        _CACHE_FILE.write_text(str(jar_path), encoding="utf-8")
    except Exception:
        pass


def _decode(b):
    """Decode captured bytes: try utf-8 first, then GBK (Chinese Windows Java)."""
    if not b:
        return ""
    for enc in ("utf-8", "gbk", "cp936"):
        try:
            return b.decode(enc)
        except Exception:
            continue
    return b.decode("utf-8", errors="replace")


def _find_near_source(source_path):
    """
    Walk up from the input file's directory looking for the project JAR.

    Inventory files usually live INSIDE the Spring Boot demo project
    (e.g. <project>/src/main/resources/static/...), whose JAR sits at
    <project>/target/. Checking the source's ancestor chain finds it without
    any broad filesystem scan.
    """
    if not source_path:
        return None
    p = normalize_path(source_path).parent
    for parent in [p, *p.parents]:
        for cand in (
            parent / "target" / JAR_NAME,
            parent / "demo" / "target" / JAR_NAME,
            parent / JAR_NAME,
        ):
            if cand.exists():
                return cand.resolve()
        tgt = parent / "target"
        if tgt.is_dir():
            try:
                for j in sorted(tgt.glob("*.jar")):
                    if "demo" in j.name.lower() or "snapshot" in j.name.lower():
                        return j.resolve()
            except Exception:
                pass
    return None


def _find_known_locations():
    """
    Safe well-known locations: skill dir / CWD / home — explicit paths only,
    plus ONE level of subdirectories under the skill dir and CWD.

    NOTE: user home subdirectories are deliberately NOT scanned. Iterating
    ~/.ssh, ~/.gnupg, ... trips sandbox permission blocks on Windows and
    slows discovery; the source-relative walk covers the real-world case.
    """
    bases = [
        Path(__file__).resolve().parent,          # skill/scripts
        Path(__file__).resolve().parent.parent,   # skill root
        Path.cwd(),
        Path.home(),
    ]
    candidates = []
    for base in bases:
        candidates.append(base / JAR_NAME)
        candidates.append(base / "demo" / "target" / JAR_NAME)
    # Subdirectory scan limited to skill dir + CWD (bases[:3]); home skipped.
    for base in bases[:3]:
        try:
            for child in base.iterdir():
                if child.is_dir():
                    candidates.append(child / JAR_NAME)
                    candidates.append(child / "demo" / "target" / JAR_NAME)
        except Exception:
            pass

    seen = set()
    for c in candidates:
        c = c.resolve()
        if c in seen:
            continue
        seen.add(c)
        if c.exists():
            return c
    return None


def find_jar(explicit=None, near_source=None):
    """
    Locate the demo JAR — JAR-FIRST by design.

    Priority:
      1. explicit path (--jar) — missing explicit path fails loud (returns None)
      2. cached last-known-good path (~/.workbuddy/excel-validate-stats.jar.txt)
      3. source-relative: walk up from the input file's dir, checking
         target/ and demo/target/ (files usually live inside the demo project)
      4. safe well-known locations (skill dir / CWD / home explicit paths,
         plus one subdir level of skill dir and CWD — NEVER home subdirs)

    Any successfully found path is cached for the next run.
    Returns a resolved Path or None (None => caller falls back to Python).
    """
    if explicit:
        p = normalize_path(explicit)
        if p.exists():
            _write_cache(p.resolve())
            print(f"[Java] 使用 JAR: {p.resolve()}  (来源: --jar 显式指定)")
            return p.resolve()
        print(f"警告: 指定的 JAR 不存在: {p}", file=sys.stderr)
        return None

    cached = _read_cache()
    if cached:
        print(f"[Java] 使用 JAR: {cached}  (来源: 上次缓存)")
        return cached

    near = _find_near_source(near_source)
    if near:
        _write_cache(near)
        print(f"[Java] 使用 JAR: {near}  (来源: 源文件上级目录)")
        return near

    known = _find_known_locations()
    if known:
        _write_cache(known)
        print(f"[Java] 使用 JAR: {known}  (来源: 已知位置)")
        return known

    print("警告: 未找到 demo-0.0.1-SNAPSHOT.jar，回退到纯 Python 实现", file=sys.stderr)
    return None


def _run_java(cmd):
    """Run a java command, decode output, raise on non-zero exit."""
    print("调用 Java 工具:", " ".join(cmd))
    try:
        res = subprocess.run(cmd, capture_output=True)
    except FileNotFoundError:
        raise RuntimeError("未找到 java 命令，请确认已安装 JDK/JRE 并加入 PATH")
    out = _decode(res.stdout).strip()
    err = _decode(res.stderr).strip()
    if res.returncode != 0:
        detail = (err or out or "").strip()
        raise RuntimeError(f"Java 执行失败 (exit {res.returncode}): {detail}")
    if out:
        print(out)


def java_validate(jar_path, source_path, overwrite=False):
    """
    Run the JAR `validate` mode. Returns the path to the produced
    <source>_重复行合并.xlsx (next to the source file).

    NOTE: the JAR's validate semantics differ from the Python engine. When it
    finds STRICT duplicate rows (same key columns AND same 进价) it MARKS them
    (重复标记 = "重复行校验不通过...未执行合并") and does NOT merge. Callers must
    check the output with find_unmerged_duplicate_rows() and fall back to the
    Python merge when marked rows exist, so the workflow still delivers a clean
    deduplicated file.
    """
    src = normalize_path(source_path)
    out_name = insert_suffix(src.name, "_重复行合并")
    out_path = src.parent / out_name
    if out_path.exists() and not overwrite:
        print(f"跳过: 已存在合并文件 {out_name}（使用 --overwrite 覆盖）")
        return out_path

    src_fp = file_fingerprint(src)  # 红线原则: 记录输入指纹, 结束后必须未变
    tmp = Path(tempfile.mkdtemp(prefix="inv_val_"))
    try:
        shutil.copy(src, tmp / src.name)
        cmd = ["java", "-jar", str(jar_path), "validate", str(tmp), src.name]
        _run_java(cmd)
        produced = tmp / out_name
        if not produced.exists():
            raise RuntimeError(f"Java validate 未生成预期文件: {out_name}")
        shutil.move(str(produced), str(out_path))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    if not verify_inputs_untouched([(src, src_fp)], ["输入文件"]):
        raise RuntimeError("红线原则违反: 输入文件在 Java validate 后被修改, 已终止")
    return out_path


def find_unmerged_duplicate_rows(xlsx_path):
    """
    Count rows in a JAR-validate output whose 重复标记 column is non-empty
    (i.e. "重复行校验不通过...未执行合并"). Returns (marked_count, first_markers).
    Used to detect when the JAR refused to merge strict duplicates so the
    caller can fall back to the Python merge engine.
    """
    import openpyxl

    p = normalize_path(xlsx_path)
    try:
        wb = openpyxl.load_workbook(p, read_only=True)
        ws = wb.active
        hdr = [c.value for c in next(ws.iter_rows())]
    except Exception:
        return 0, []
    mark_col = None
    for i, h in enumerate(hdr):
        if h and "重复标记" in str(h):
            mark_col = i
            break
    if mark_col is None:
        wb.close()
        return 0, []
    marked, markers = 0, []
    try:
        for row in ws.iter_rows(min_row=2, values_only=True):
            if mark_col < len(row) and row[mark_col] is not None and str(row[mark_col]).strip() != "":
                marked += 1
                if len(markers) < 3:
                    markers.append(str(row[mark_col]).strip())
    finally:
        wb.close()
    return marked, markers


def java_compare(jar_path, source_path, compare_path, out_dir, overwrite=False):
    """
    Run the JAR `compare` mode. Returns the path to the produced
    <source>_差异.xlsx (in out_dir).
    """
    src = normalize_path(source_path)
    cmp = normalize_path(compare_path)
    out_dir = normalize_path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_name = insert_suffix(src.name, "_差异")
    diff_path = out_dir / out_name
    if diff_path.exists() and not overwrite:
        print(f"跳过: 已存在差异文件 {out_name}（使用 --overwrite 覆盖）")
        return diff_path

    if src.name == cmp.name:
        raise RuntimeError("两个输入文件不能同名")

    # 红线原则: 记录两个输入文件的指纹, 结束后必须均未变
    src_fp = file_fingerprint(src)
    cmp_fp = file_fingerprint(cmp)

    tmp = Path(tempfile.mkdtemp(prefix="inv_cmp_"))
    try:
        shutil.copy(src, tmp / src.name)
        shutil.copy(cmp, tmp / cmp.name)
        cmd = ["java", "-jar", str(jar_path), "compare", str(tmp), src.name, cmp.name]
        _run_java(cmd)
        produced = tmp / out_name
        if not produced.exists():
            raise RuntimeError(f"Java compare 未生成预期文件: {out_name}")
        shutil.move(str(produced), str(diff_path))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    if not verify_inputs_untouched(
        [(src, src_fp), (cmp, cmp_fp)], ["文件A(源文件)", "文件B(对比文件)"]
    ):
        raise RuntimeError("红线原则违反: 输入文件在 Java compare 后被修改, 已终止")
    return diff_path
