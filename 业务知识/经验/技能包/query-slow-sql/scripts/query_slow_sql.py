#!/usr/bin/env python3
"""Query slow SQL logs from MySQL slow_sql_log table by time range.

Exports the RAW records (all table fields, no aggregation/statistics)
as an Excel (.xlsx) file inside the user-specified workspace directory.
"""

import argparse
import datetime
import decimal
import os
import sys

import pymysql
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --- Database connection config ---
DB_HOST = "rm-bp1151j106lak7koawo.mysql.rds.aliyuncs.com"
DB_PORT = 3306
DB_USER = "slow_sql_user"
DB_PASSWORD = "Bn&KITDHz6"
DB_NAME = "jy_monitor"  # slow_sql_log 表所在的数据库


def parse_args():
    parser = argparse.ArgumentParser(
        description="Query slow SQL logs by time range and database, export raw data to Excel"
    )
    parser.add_argument(
        "--start-time", required=False,
        help="开始时间，格式 YYYY-MM-DD 或 YYYY-MM-DD HH:MM:SS（未指定时默认今天 00:00:00）"
    )
    parser.add_argument(
        "--end-time", required=False,
        help="结束时间，格式 YYYY-MM-DD 或 YYYY-MM-DD HH:MM:SS（未指定时默认今天 23:59:59）"
    )
    parser.add_argument(
        "--database", required=False,
        help="数据库名（未指定时可先执行 --list-databases 查看可选数据库列表）"
    )
    parser.add_argument(
        "--workspace", required=False,
        help="工作空间（项目源码绝对路径），例如 D:\\projects\\jy-backend；必须由用户明确指定，禁止使用默认工作空间"
    )
    parser.add_argument(
        "--list-databases", action="store_true",
        help="仅查询 slow_sql_log 表中的数据库列表（SELECT DISTINCT db_name FROM slow_sql_log）并打印，用于用户选择数据库"
    )
    return parser.parse_args()


def validate_args(args):
    """Validate that required arguments are provided.

    start-time / end-time are optional: if omitted, they default to today
    (00:00:00 ~ 23:59:59) via resolve_time_range(). database and workspace
    are still mandatory for the main query. workspace must be explicitly
    provided by the user — never substitute the default/current directory.
    """
    errors = []
    if not args.database:
        errors.append("缺少参数：--database（数据库名，可先运行 --list-databases 查看可选列表）")
    if not args.workspace:
        errors.append("缺少参数：--workspace（工作空间/项目源码绝对路径，必须由用户明确指定，禁止使用默认工作空间）")
    if errors:
        print("参数不完整。请指定数据库名和工作空间；时间未指定时默认查询今天。")
        for e in errors:
            print(f"  - {e}")
        print("若不确定数据库名，可先运行 --list-databases 查看可选数据库列表。")
        print("工作空间必须由用户明确指定（导出慢SQL的目录），禁止使用 WorkBuddy 默认工作空间。")
        print(
            '示例：python query_slow_sql.py --start-time "2026-07-27 00:00:00" '
            '--end-time "2026-07-27 23:59:59" --database "mydb" '
            '--workspace "D:\\projects\\jy-backend"'
        )
        sys.exit(1)


def resolve_time_range(args):
    """Resolve the time range; if not specified, default to today.

    Returns (start_time, end_time) as strings in 'YYYY-MM-DD HH:MM:SS' format.
    """
    today = datetime.date.today()
    start_time = args.start_time if args.start_time else f"{today} 00:00:00"
    end_time = args.end_time if args.end_time else f"{today} 23:59:59"
    return start_time, end_time


def to_excel_value(v):
    """Convert a DB value into something openpyxl can write as-is."""
    if v is None:
        return None
    if isinstance(v, decimal.Decimal):
        return float(v)
    if isinstance(v, bytes):
        return v.decode("utf-8", errors="replace")
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return v
    return v


def query_slow_sql(start_time, end_time, database_name, workspace):
    """Connect to MySQL, query slow_sql_log, and export raw rows to Excel."""
    connection = None
    try:
        connection = pymysql.connect(
            host=DB_HOST,
            port=DB_PORT,
            user=DB_USER,
            password=DB_PASSWORD,
            database=DB_NAME,
            charset="utf8mb4",
            connect_timeout=10,
            read_timeout=30,
        )

        # Fetch ALL fields — raw data, no projection
        sql = """
            SELECT *
            FROM slow_sql_log
            WHERE execution_start_time >= %s
              AND execution_start_time <= %s
              AND db_name = %s
            ORDER BY query_time DESC
        """

        with connection.cursor(pymysql.cursors.DictCursor) as cursor:
            cursor.execute(sql, (start_time, end_time, database_name))
            rows = cursor.fetchall()

        if not rows:
            print(f"未查询到 {start_time} 至 {end_time} 数据库 {database_name} 的慢SQL记录，未生成 Excel 文件。")
            return

        # Column order follows the raw table layout (from the first row)
        columns = list(rows[0].keys())

        # Output file lives inside the workspace
        safe_start = start_time.replace(":", "").replace(" ", "_")
        safe_end = end_time.replace(":", "").replace(" ", "_")
        filename = f"slow_sql_{safe_start}_to_{safe_end}_{database_name}.xlsx"
        output_path = os.path.join(workspace, filename)

        wb = Workbook()
        ws = wb.active
        ws.title = "慢SQL原始数据"

        # Header row
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        ws.append(columns)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Raw data rows — NO aggregation, NO statistics, NO summary
        for row in rows:
            ws.append([to_excel_value(row.get(c)) for c in columns])

        # Column widths
        for idx, col in enumerate(columns, start=1):
            letter = get_column_letter(idx)
            max_len = len(str(col))
            for cell in ws[letter]:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[letter].width = min(max_len + 2, 60)

        ws.freeze_panes = "A2"
        wb.save(output_path)

        print(f"慢SQL原始数据已导出: {output_path}")

    except pymysql.MySQLError as e:
        print(f"数据库连接或查询失败: {e}", file=sys.stderr)
        sys.exit(1)
    except OSError as e:
        print(f"写入 Excel 文件失败（请确认工作空间目录存在且可写）: {e}", file=sys.stderr)
        sys.exit(1)
    finally:
        if connection:
            connection.close()


def list_databases():
    """Query SELECT DISTINCT db_name FROM slow_sql_log and print the database list.

    Used when the user did not specify a database: the agent shows this list
    and asks the user to pick one before running the main query.
    """
    connection = None
    try:
        connection = pymysql.connect(
            host=DB_HOST,
            port=DB_PORT,
            user=DB_USER,
            password=DB_PASSWORD,
            database=DB_NAME,
            charset="utf8mb4",
            connect_timeout=10,
            read_timeout=30,
        )

        with connection.cursor() as cursor:
            cursor.execute("SELECT DISTINCT db_name FROM slow_sql_log ORDER BY db_name")
            rows = cursor.fetchall()

        db_names = [row[0] for row in rows]
        if not db_names:
            print("slow_sql_log 表中暂无数据库记录。")
            return

        print("slow_sql_log 表中的数据库列表：")
        for i, name in enumerate(db_names, start=1):
            print(f"{i}. {name}")

    except pymysql.MySQLError as e:
        print(f"数据库连接或查询失败: {e}", file=sys.stderr)
        sys.exit(1)
    finally:
        if connection:
            connection.close()


def main():
    args = parse_args()
    if args.list_databases:
        list_databases()
        return
    validate_args(args)
    start_time, end_time = resolve_time_range(args)
    query_slow_sql(start_time, end_time, args.database, args.workspace)


if __name__ == "__main__":
    main()
